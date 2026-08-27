"""
Inspección SEGURA del esquema de la entrega IMARPE — PredictaMAR Litoral

QUÉ HACE Y QUÉ NO:
Describe la FORMA del Excel entregado por IMARPE -- hojas, tamaños,
encabezados, tipos, faltantes, rango de fechas, presencia de hora y conteos de
unicidad -- y NUNCA su CONTENIDO. No imprime filas, ni coordenadas concretas,
ni nombres de personas o embarcaciones, ni valores observados. Esa asimetría
existe porque la entrega está sujeta a restricciones de transferencia: el
esquema puede discutirse y versionarse; los datos, no.

MINIMIZACIÓN DE ACCESO:
Los nombres de hoja se listan SIN cargar su contenido, y solo las hojas
requeridas se materializan. Ninguna otra hoja se abre ni se recorre. Es una
obligación de tratamiento, no una optimización: si el libro contuviera una
hoja con datos que no nos corresponden, no debemos ni leerla.

DÓNDE VIVEN LOS DATOS:
Fuera del repositorio, siempre. La ruta llega EXCLUSIVAMENTE por la variable de
entorno declarada en la configuración. No hay ninguna ruta real en este archivo
ni en el YAML, y no debe haberla nunca.

SIN RED:
No realiza ninguna petición. Lee un archivo local y nada más.

ENCABEZADOS EXACTOS:
Se leen de la fila de encabezado tal como están, sin normalizar, sin recortar y
sin añadir sufijos. En particular NO se confía en DataFrame.columns, porque
pandas desambigua los duplicados renombrándolos: un libro con dos columnas
iguales produciría encabezados que no existen en el archivo. Encabezados
vacíos, duplicados o no interpretables detienen la inspección con error.

NO SE ADIVINAN ROLES:
Qué columna es la fecha, la hora, las coordenadas, la estación o la profundidad
se toma EXCLUSIVAMENTE de column_roles. Si un rol no está declarado, la métrica
que depende de él se reporta como NO RESUELTA y se listan los candidatos
observados -- candidatos, no elecciones. En particular, que exista una sola
columna de tipo fecha-hora no la convierte en la fecha oficial, y declarar
'fecha' no dice nada sobre la hora: si la hora viaja en esa misma columna, hay
que declarar ese mismo encabezado también en '.hora'.

Un rol declarado cuya columna NO existe es un error, no una degradación
silenciosa: significa que la configuración y el archivo discrepan, y seguir
adelante produciría métricas calculadas sobre otra cosa.

CUARENTENA, NO CORRECCIÓN:
Las fechas fuera del periodo esperado se CUENTAN y se marcan. Las no
interpretables también se cuentan, igual que las profundidades no numéricas.
Ninguna se corrige, se infiere ni se reasigna: una fecha anómala puede ser el
síntoma de un error de exportación con más registros afectados, y corregirla
borraría la evidencia.

FLUORESCENCIA:
Se declara NO CUANTITATIVA en el informe. No es clorofila calibrada y no puede
validar el producto de clorofila sin documentación de calibración.

FALTANTES:
Permanecen faltantes. Este módulo no rellena, no interpola y no convierte a
cero en ningún camino.
"""

from __future__ import annotations

import datetime as _dt
import os
import subprocess
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

import pandas as pd
import yaml

DEFAULT_CONFIG_PATH = "config/imarpe_source.yaml"

# Raiz del repositorio, deducida de la UBICACION DE ESTE ARCHIVO y no del
# directorio de trabajo. Usar "." haria que la compuerta dependiera de donde se
# lance el comando: ejecutando desde un subdirectorio, un Excel situado en otra
# carpeta del repositorio quedaria "fuera" y la comprobacion no dispararia.
REPO_ROOT = Path(__file__).resolve().parents[1]

# Extensiones portadoras de datos. Respaldo si el YAML no las declara.
DEFAULT_RESTRICTED_SUFFIXES = (".zip", ".xlsx", ".xls", ".csv", ".pdf")

FLUORESCENCE_DECLARATION = (
    "NO CUANTITATIVA: la fluorescencia no es clorofila calibrada y no puede "
    "usarse para validar el producto de clorofila mientras no exista "
    "documentacion de calibracion."
)

MISSING_POLICY = (
    "Los faltantes permanecen faltantes: sin relleno, sin interpolacion y sin "
    "conversion a cero. Las fechas no interpretables se cuentan; sus valores no "
    "se muestran ni se corrigen."
)

QUARANTINE_POLICY = (
    "Las fechas fuera del periodo esperado se marcan para CUARENTENA y nunca "
    "se corrigen ni se infieren."
)

GITIGNORE_POLICY = (
    ".gitignore evita versionar los datos, pero NO autoriza copiarlos a "
    "Codespaces ni a ningun entorno remoto: son dos controles distintos."
)

# Tipos aceptados en una celda de encabezado. Cualquier otra cosa se considera
# no interpretable y detiene la inspeccion.
_HEADER_SCALARS = (str, int, float, Decimal, _dt.datetime, _dt.date, _dt.time)


class ImarpeInspectionError(RuntimeError):
    """Error de inspección de la fuente IMARPE."""


class SourceNotDeclaredError(ImarpeInspectionError):
    """La variable de entorno con la ruta no está declarada o está vacía."""


class SourceNotFoundError(ImarpeInspectionError):
    """La ruta declarada no existe o no es un archivo legible."""


class MissingSheetError(ImarpeInspectionError):
    """Falta una hoja requerida en el libro."""


class UnreadableSchemaError(ImarpeInspectionError):
    """El libro, una hoja o un encabezado no son interpretables como esquema."""


class GovernanceCheckError(ImarpeInspectionError):
    """La comprobación de datos restringidos no pudo completarse o falló."""


class ExecutionEnvironmentError(ImarpeInspectionError):
    """El entorno o la ubicación del archivo no autorizan ejecutar la inspección."""


@dataclass(frozen=True)
class SheetPayload:
    """
    Material mínimo de una hoja: la fila de encabezado EXACTA y los datos.

    Los encabezados viajan aparte del marco a propósito: son la verdad del
    archivo, mientras que las columnas del marco ya pasaron por pandas.
    """

    raw_headers: tuple[Any, ...]
    frame: pd.DataFrame


@dataclass(frozen=True)
class SheetSchema:
    """Esquema de una hoja. Contiene descripciones, nunca contenido."""

    name: str
    n_rows: int
    n_cols: int
    headers: tuple[str, ...]
    dtypes: Mapping[str, str]
    missing_counts: Mapping[str, int]
    # Candidatos OBSERVADOS, no elecciones. Informativos.
    datetime_candidate_headers: tuple[str, ...]
    # Metricas oficiales: solo con rol declarado.
    date_role: str | None
    time_role: str | None
    depth_role: str | None
    time_source: str | None
    has_time_component: bool | None
    date_min: str | None
    date_max: str | None
    n_out_of_period: int | None
    n_unparseable_dates: int | None
    n_unparseable_times: int | None
    fluorescence_headers: tuple[str, ...]
    n_unique_coordinates: int | None
    n_unique_stations: int | None
    n_unique_depths: int | None
    n_non_numeric_depths: int | None
    unresolved_metrics: tuple[str, ...] = ()


@dataclass(frozen=True)
class InspectionReport:
    """Informe completo. Solo descripciones de esquema."""

    source_env_var: str
    sheet_names: tuple[str, ...]
    materialized_sheets: tuple[str, ...]
    sheets: tuple[SheetSchema, ...]
    period_start: str
    period_end: str
    fluorescence_declaration: str = FLUORESCENCE_DECLARATION
    missing_policy: str = MISSING_POLICY
    quarantine_policy: str = QUARANTINE_POLICY
    gitignore_policy: str = GITIGNORE_POLICY


# ---------------------------------------------------------------------------
# Configuración y resolución de la fuente
# ---------------------------------------------------------------------------
def load_config(config_path: str | Path = DEFAULT_CONFIG_PATH) -> dict:
    """Carga la declaración de la fuente. No contiene rutas ni datos."""
    ruta = Path(config_path)
    if not ruta.is_file():
        raise ImarpeInspectionError(
            f"No se encuentra la declaracion de la fuente en {ruta}. "
            "Sin ella no se conocen las hojas requeridas ni la politica de cuarentena."
        )
    try:
        cfg = yaml.safe_load(ruta.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ImarpeInspectionError(
            f"La declaracion de la fuente {ruta} no es YAML valido: {type(exc).__name__}"
        ) from exc
    if not isinstance(cfg, dict) or "fuente" not in cfg or "hojas_requeridas" not in cfg:
        raise ImarpeInspectionError(
            f"La declaracion de la fuente {ruta} no tiene la estructura esperada "
            "(se requieren al menos las claves 'fuente' y 'hojas_requeridas')."
        )
    return cfg


def resolve_source_path(cfg: Mapping, env: Mapping[str, str] | None = None) -> Path:
    """
    Resuelve la ruta del Excel EXCLUSIVAMENTE desde la variable de entorno
    declarada en la configuración.

    No acepta argumento de ruta, no busca en ubicaciones por defecto y no tiene
    ningun valor de reserva. Esa rigidez es intencional: cualquier camino
    alternativo terminaria, tarde o temprano, con una ruta real en el
    repositorio.
    """
    nombre = cfg["fuente"]["variable_entorno"]
    entorno = os.environ if env is None else env
    valor = (entorno.get(nombre) or "").strip()
    if not valor:
        raise SourceNotDeclaredError(
            f"La variable de entorno {nombre} no esta definida. La ruta del Excel "
            "de IMARPE solo puede llegar por esa via: no se escribe en el codigo "
            "ni en la configuracion."
        )
    ruta = Path(valor).expanduser()
    if not ruta.is_file():
        raise SourceNotFoundError(
            f"La ruta declarada en {nombre} no apunta a un archivo legible. "
            "Se omite la ruta de este mensaje a proposito."
        )
    return ruta


# ---------------------------------------------------------------------------
# Capa de acceso (aislada y sustituible en pruebas)
# ---------------------------------------------------------------------------
def list_sheet_names(path: Path) -> tuple[str, ...]:
    """
    Nombres de hoja SIN cargar contenido.

    openpyxl en modo solo lectura expone la lista de hojas a partir de los
    metadatos del libro: no recorre celdas ni materializa filas.
    """
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return tuple(str(n) for n in wb.sheetnames)
        finally:
            wb.close()
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ImarpeInspectionError(
            "Falta la dependencia openpyxl, necesaria para leer .xlsx."
        ) from exc
    except Exception as exc:
        raise UnreadableSchemaError(
            f"No se pudo listar las hojas del libro: {type(exc).__name__}. "
            "Revise que sea un .xlsx legible y no un archivo protegido o corrupto."
        ) from exc


def read_sheets(path: Path, sheets: Sequence[str]) -> dict[str, SheetPayload]:
    """
    Materializa ÚNICAMENTE las hojas indicadas.

    De cada una toma dos cosas: la fila de encabezado EXACTA, leida con
    openpyxl sin pasar por la desambiguacion de pandas, y los datos, leidos con
    pandas para poder inferir tipos y contar faltantes. Ninguna hoja fuera de
    `sheets` se abre.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ImarpeInspectionError(
            "Falta la dependencia openpyxl, necesaria para leer .xlsx."
        ) from exc

    salida: dict[str, SheetPayload] = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise UnreadableSchemaError(
            f"El libro no pudo abrirse: {type(exc).__name__}."
        ) from exc
    try:
        libro = pd.ExcelFile(path, engine="openpyxl")
    except Exception as exc:
        wb.close()
        raise UnreadableSchemaError(
            f"El libro no pudo leerse como esquema tabular: {type(exc).__name__}."
        ) from exc

    try:
        for nombre in sheets:
            try:
                ws = wb[nombre]
                fila = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            except Exception as exc:
                raise UnreadableSchemaError(
                    f"No se pudo leer la fila de encabezado de la hoja {nombre!r}: "
                    f"{type(exc).__name__}."
                ) from exc
            try:
                marco = libro.parse(sheet_name=nombre)
            except Exception as exc:
                raise UnreadableSchemaError(
                    f"La hoja {nombre!r} no pudo leerse como tabla: {type(exc).__name__}."
                ) from exc
            salida[nombre] = SheetPayload(raw_headers=tuple(fila or ()), frame=marco)
    finally:
        wb.close()
        try:
            libro.close()
        except Exception:
            pass
    return salida


# ---------------------------------------------------------------------------
# Encabezados exactos
# ---------------------------------------------------------------------------
def exact_headers(sheet: str, raw: Sequence[Any]) -> tuple[str, ...]:
    """
    Valida y devuelve los encabezados EXACTOS de una hoja.

    No normaliza, no recorta espacios y no añade sufijos. Falla ante:
      - fila de encabezado ausente o vacia,
      - celdas vacias o en blanco,
      - celdas no interpretables como texto,
      - encabezados duplicados.

    El duplicado es el caso que motiva no usar DataFrame.columns: pandas lo
    resuelve renombrando la segunda aparicion, de modo que el informe mostraria
    un encabezado que no existe en el archivo.
    """
    if not raw:
        raise UnreadableSchemaError(
            f"La hoja {sheet!r} no tiene fila de encabezado legible."
        )

    nombres: list[str] = []
    for i, celda in enumerate(raw, start=1):
        if celda is None:
            raise UnreadableSchemaError(
                f"La hoja {sheet!r} tiene un encabezado vacio en la posicion {i}. "
                "No se rellena ni se renombra: corrija la fuente."
            )
        if not isinstance(celda, _HEADER_SCALARS):
            raise UnreadableSchemaError(
                f"La hoja {sheet!r} tiene un encabezado no interpretable en la "
                f"posicion {i} (tipo {type(celda).__name__})."
            )
        texto = celda if isinstance(celda, str) else str(celda)
        if not texto.strip():
            raise UnreadableSchemaError(
                f"La hoja {sheet!r} tiene un encabezado en blanco en la posicion {i}."
            )
        nombres.append(texto)  # exacto: sin strip, sin normalizar

    vistos: dict[str, int] = {}
    duplicados: list[str] = []
    for n in nombres:
        vistos[n] = vistos.get(n, 0) + 1
        if vistos[n] == 2:
            duplicados.append(n)
    if duplicados:
        raise UnreadableSchemaError(
            f"La hoja {sheet!r} tiene encabezados duplicados: {duplicados}. "
            "No se renombran ni se les añade sufijo: corrija la fuente o declare "
            "el criterio de desambiguacion."
        )
    return tuple(nombres)


# ---------------------------------------------------------------------------
# Utilidades de esquema
# ---------------------------------------------------------------------------
def _period_bounds(cfg: Mapping) -> tuple[_dt.date, _dt.date]:
    p = cfg.get("periodo_esperado") or {}
    try:
        inicio = pd.Timestamp(p["inicio"]).date()
        fin = pd.Timestamp(p["fin"]).date()
    except Exception as exc:
        raise ImarpeInspectionError(
            "El periodo esperado de la configuracion no es interpretable como fechas."
        ) from exc
    return inicio, fin


def _declared_role(cfg: Mapping, sheet: str, rol: str) -> str | None:
    roles = (cfg.get("column_roles") or {}).get(sheet) or {}
    valor = roles.get(rol)
    if valor is None:
        return None
    texto = str(valor)
    return texto if texto.strip() else None


def _resolve_role(cfg: Mapping, sheet: str, rol: str, headers: Sequence[str]) -> str | None:
    """
    Devuelve el encabezado declarado para un rol, o None si no esta declarado.

    Si esta declarado pero NO existe en la hoja, es un error: configuracion y
    archivo discrepan, y calcular sobre otra columna seria peor que no calcular.
    """
    declarado = _declared_role(cfg, sheet, rol)
    if declarado is None:
        return None
    if declarado not in headers:
        raise UnreadableSchemaError(
            f"La hoja {sheet!r} no contiene la columna {declarado!r} declarada como "
            f"rol {rol!r} en column_roles. No se busca una columna parecida ni se "
            "continua sin ella."
        )
    return declarado


def _column(frame: pd.DataFrame, headers: Sequence[str], nombre: str) -> pd.Series:
    """Serie por POSICION del encabezado exacto, sin depender de frame.columns."""
    return frame.iloc[:, headers.index(nombre)]


def _datetime_candidates(frame: pd.DataFrame, headers: Sequence[str]) -> tuple[str, ...]:
    """
    Encabezados cuya columna quedo tipada como fecha-hora tras la lectura.

    Son CANDIDATOS observados, nunca una eleccion: se informan para que la
    Fase 1B declare el rol con conocimiento, no para sustituirlo.
    """
    salida = []
    for i, h in enumerate(headers):
        if i < frame.shape[1] and pd.api.types.is_datetime64_any_dtype(frame.iloc[:, i]):
            salida.append(h)
    return tuple(salida)


def _clock_parts(valor: Any) -> tuple[int, int, int] | None:
    """
    (hora, minuto, segundo) de un valor, o None si no es interpretable.

    Se resuelve valor a valor y no de forma vectorizada a proposito: una
    columna de hora puede traer datetime.time, marcas completas o texto, y el
    parseo vectorizado de pandas mezcla esos casos con una advertencia y un
    formato adivinado. Aqui cada tipo se trata explicitamente y lo que no
    encaja se cuenta como no interpretable en vez de forzarse.
    """
    if valor is None:
        return None
    if isinstance(valor, _dt.time):
        return (valor.hour, valor.minute, valor.second)
    if isinstance(valor, _dt.datetime):
        return (valor.hour, valor.minute, valor.second)
    if isinstance(valor, _dt.date):
        return (0, 0, 0)
    try:
        marca = pd.Timestamp(valor)
    except Exception:
        return None
    if marca is pd.NaT or pd.isna(marca):
        return None
    return (int(marca.hour), int(marca.minute), int(marca.second))


def _time_presence(serie: pd.Series) -> tuple[bool, int]:
    """
    (hay_hora, n_no_interpretables) para una columna que deberia portar hora.

    Todo a medianoche significa, en la practica, que la fuente no registro la
    hora: es el caso que impide emparejar con un producto de cuatro instantes
    diarios, asi que conviene detectarlo desde el esquema.
    """
    presentes = serie[serie.notna()]
    if presentes.empty:
        return False, 0

    hay = False
    malos = 0
    for valor in presentes:
        partes = _clock_parts(valor)
        if partes is None:
            malos += 1
            continue
        if partes != (0, 0, 0):
            hay = True
    return hay, malos


def _fluorescence_headers(headers: Sequence[str], tokens: Sequence[str]) -> tuple[str, ...]:
    """
    Encabezados REALES que coinciden con alguno de los tokens declarados.

    No inventa un nombre de columna: informa cuales de los encabezados que ya
    existen coinciden, para poder marcarlos como no cuantitativos.
    """
    bajos = [str(t).lower() for t in tokens]
    return tuple(h for h in headers if any(t in h.lower() for t in bajos))


# ---------------------------------------------------------------------------
# Inspección de una hoja
# ---------------------------------------------------------------------------
def inspect_sheet(name: str, payload: SheetPayload, cfg: Mapping) -> SheetSchema:
    """
    Describe una hoja sin exponer su contenido.

    Las metricas oficiales -- rango de fechas, cuarentena, hora, coordenadas,
    estaciones y profundidad -- se calculan SOLO con el rol declarado. Sin rol,
    se reportan como no resueltas junto con los candidatos observados, y esa
    lista es justamente el insumo de la Fase 1B.
    """
    marco = payload.frame
    if not isinstance(marco, pd.DataFrame):
        raise UnreadableSchemaError(f"La hoja {name!r} no pudo leerse como tabla.")

    headers = exact_headers(name, payload.raw_headers)
    if marco.shape[1] != len(headers):
        raise UnreadableSchemaError(
            f"La hoja {name!r} declara {len(headers)} encabezados pero el cuerpo "
            f"tiene {marco.shape[1]} columnas. No se ajusta ninguno de los dos."
        )

    dtypes = {h: str(marco.iloc[:, i].dtype) for i, h in enumerate(headers)}
    missing = {h: int(marco.iloc[:, i].isna().sum()) for i, h in enumerate(headers)}
    candidatos = _datetime_candidates(marco, headers)

    inicio, fin = _period_bounds(cfg)
    sin_resolver: list[str] = []

    # ---- rol fecha ----
    rol_fecha = _resolve_role(cfg, name, "fecha", headers)
    if rol_fecha is None:
        fecha_min = fecha_max = None
        fuera = None
        n_fechas_malas = None
        sin_resolver.append(
            "rango_de_fechas y fechas_fuera_de_periodo: falta declarar "
            f"column_roles.{name}.fecha (candidatos fecha-hora observados: "
            f"{list(candidatos) or 'ninguno'})"
        )
        serie_fecha = None
    else:
        serie_fecha = _column(marco, headers, rol_fecha)
        presentes = serie_fecha[serie_fecha.notna()]
        convertidos = pd.to_datetime(presentes, errors="coerce")
        n_fechas_malas = int(convertidos.isna().sum())
        validos = convertidos.dropna()
        if validos.empty:
            fecha_min = fecha_max = None
            fuera = 0
        else:
            fechas = validos.dt.date
            fecha_min = str(fechas.min())
            fecha_max = str(fechas.max())
            fuera = int(((fechas < inicio) | (fechas > fin)).sum())

    # ---- rol hora: SIEMPRE explicito, sin heredar de fecha ----
    # Si la hora viaja dentro de la misma columna que la fecha, hay que
    # declarar ESE MISMO encabezado en los dos roles. Deducirla de 'fecha'
    # cuando 'hora' esta en null seria otro fallback: convertiria la ausencia
    # de una declaracion en una afirmacion sobre el archivo.
    rol_hora = _resolve_role(cfg, name, "hora", headers)
    if rol_hora is not None:
        serie_hora = _column(marco, headers, rol_hora)
        hay_hora, n_horas_malas = _time_presence(serie_hora)
        origen = ("misma_columna_que_fecha" if rol_hora == rol_fecha
                  else "columna_hora_separada")
    else:
        hay_hora = None
        n_horas_malas = None
        origen = None
        sin_resolver.append(
            f"presencia_de_hora: falta declarar column_roles.{name}.hora. Si la "
            f"hora viaja en la misma columna que la fecha, declare ese mismo "
            f"encabezado tambien en .hora; candidatos fecha-hora observados: "
            f"{list(candidatos) or 'ninguno'}"
        )

    # ---- coordenadas ----
    rol_lat = _resolve_role(cfg, name, "latitud", headers)
    rol_lon = _resolve_role(cfg, name, "longitud", headers)
    if rol_lat is not None and rol_lon is not None:
        pares = pd.DataFrame({
            "a": _column(marco, headers, rol_lat),
            "b": _column(marco, headers, rol_lon),
        }).dropna()
        n_coords = int(len(pares.drop_duplicates()))
    else:
        n_coords = None
        sin_resolver.append(
            f"coordenadas_unicas: falta declarar column_roles.{name}.latitud y .longitud"
        )

    # ---- estaciones ----
    rol_est = _resolve_role(cfg, name, "estacion", headers)
    if rol_est is not None:
        n_est = int(_column(marco, headers, rol_est).dropna().nunique())
    else:
        n_est = None
        sin_resolver.append(f"estaciones_unicas: falta declarar column_roles.{name}.estacion")

    # ---- profundidad ----
    rol_prof = _resolve_role(cfg, name, "profundidad", headers)
    if rol_prof is not None:
        serie_prof = _column(marco, headers, rol_prof)
        presentes_prof = serie_prof[serie_prof.notna()]
        numericos = pd.to_numeric(presentes_prof, errors="coerce")
        n_prof_malas = int(numericos.isna().sum())
        n_prof_unicas = int(numericos.dropna().nunique())
    else:
        n_prof_unicas = None
        n_prof_malas = None
        sin_resolver.append(
            f"niveles_de_profundidad: falta declarar column_roles.{name}.profundidad"
        )

    tokens = ((cfg.get("fluorescencia") or {}).get("tokens_encabezado")) or []
    fluo = _fluorescence_headers(headers, tokens)

    return SheetSchema(
        name=name,
        n_rows=int(len(marco)),
        n_cols=len(headers),
        headers=headers,
        dtypes=dtypes,
        missing_counts=missing,
        datetime_candidate_headers=candidatos,
        date_role=rol_fecha,
        time_role=rol_hora,
        depth_role=rol_prof,
        time_source=origen,
        has_time_component=hay_hora,
        date_min=fecha_min,
        date_max=fecha_max,
        n_out_of_period=fuera,
        n_unparseable_dates=n_fechas_malas,
        n_unparseable_times=n_horas_malas,
        fluorescence_headers=fluo,
        n_unique_coordinates=n_coords,
        n_unique_stations=n_est,
        n_unique_depths=n_prof_unicas,
        n_non_numeric_depths=n_prof_malas,
        unresolved_metrics=tuple(sin_resolver),
    )


def inspect_workbook(
    path: Path,
    cfg: Mapping,
    lister: Callable[[Path], Sequence[str]] | None = None,
    reader: Callable[[Path, Sequence[str]], Mapping[str, SheetPayload]] | None = None,
) -> InspectionReport:
    """
    Inspecciona el libro y devuelve el informe de esquema.

    Primero lista los nombres de hoja sin cargar contenido, comprueba que estén
    las requeridas y solo entonces materializa ESAS. Si falta una hoja
    requerida falla con mensaje claro: no se busca una hoja parecida, no se toma
    la primera disponible y no se continua con lo que haya.
    """
    listar = lister or list_sheet_names
    leer = reader or read_sheets

    nombres = tuple(str(n) for n in listar(path))
    requeridas = [str(h) for h in cfg["hojas_requeridas"]]
    faltan = [h for h in requeridas if h not in nombres]
    if faltan:
        raise MissingSheetError(
            "Faltan hojas requeridas en el libro: %s. Hojas presentes: %s. "
            "No se sustituyen por hojas de nombre parecido." % (faltan, list(nombres))
        )

    cargadas = leer(path, tuple(requeridas))
    ausentes = [h for h in requeridas if h not in cargadas]
    if ausentes:
        raise UnreadableSchemaError(
            "El lector no devolvio las hojas requeridas: %s." % ausentes
        )

    inicio, fin = _period_bounds(cfg)
    esquemas = tuple(inspect_sheet(h, cargadas[h], cfg) for h in requeridas)
    return InspectionReport(
        source_env_var=cfg["fuente"]["variable_entorno"],
        sheet_names=nombres,
        materialized_sheets=tuple(requeridas),
        sheets=esquemas,
        period_start=str(inicio),
        period_end=str(fin),
    )


# ---------------------------------------------------------------------------
# Formato del informe (solo descripciones)
# ---------------------------------------------------------------------------
def format_report(report: InspectionReport) -> str:
    """
    Formatea el informe. NO imprime filas, coordenadas concretas, nombres
    personales ni valores observados: solo nombres de columna, tipos, conteos,
    banderas y el rango de fechas.
    """
    out: list[str] = []
    add = out.append
    add("=== INSPECCION DE ESQUEMA IMARPE (Fase 1A) ===")
    add("ruta declarada por: %s (no se imprime)" % report.source_env_var)
    add("periodo esperado: %s a %s" % (report.period_start, report.period_end))
    add("hojas presentes en el libro: %s" % list(report.sheet_names))
    add("hojas materializadas (unicas leidas): %s" % list(report.materialized_sheets))
    for s in report.sheets:
        add("")
        add("--- hoja %s ---" % s.name)
        add("  filas=%d columnas=%d" % (s.n_rows, s.n_cols))
        add("  encabezados exactos:")
        for h in s.headers:
            add("    %-40s tipo=%-16s faltantes=%d" % (h, s.dtypes[h], s.missing_counts[h]))
        add("  candidatos fecha-hora observados: %s"
            % (list(s.datetime_candidate_headers) or "ninguno"))
        add("  rol fecha declarado: %s" % (s.date_role or "no declarado"))
        add("  rol hora declarado: %s" % (s.time_role or "no declarado"))
        add("  rol profundidad declarado: %s" % (s.depth_role or "no declarado"))
        add("  origen de la hora: %s" % (s.time_source or "no resuelto"))
        add("  hora presente: %s" %
            ("no resuelto" if s.has_time_component is None else s.has_time_component))
        add("  rango de fechas: %s a %s" % (s.date_min, s.date_max))
        add("  fechas fuera de periodo (CUARENTENA): %s" %
            ("no resuelto" if s.n_out_of_period is None else s.n_out_of_period))
        add("  fechas no interpretables: %s" %
            ("no resuelto" if s.n_unparseable_dates is None else s.n_unparseable_dates))
        add("  horas no interpretables: %s" %
            ("no resuelto" if s.n_unparseable_times is None else s.n_unparseable_times))
        add("  coordenadas unicas: %s" %
            ("no resuelto" if s.n_unique_coordinates is None else s.n_unique_coordinates))
        add("  estaciones unicas: %s" %
            ("no resuelto" if s.n_unique_stations is None else s.n_unique_stations))
        add("  niveles de profundidad unicos: %s" %
            ("no resuelto" if s.n_unique_depths is None else s.n_unique_depths))
        add("  profundidades no numericas: %s" %
            ("no resuelto" if s.n_non_numeric_depths is None else s.n_non_numeric_depths))
        if s.fluorescence_headers:
            add("  encabezados de fluorescencia: %s" % list(s.fluorescence_headers))
            add("    %s" % FLUORESCENCE_DECLARATION)
        else:
            add("  encabezados de fluorescencia: ninguno coincide con los tokens declarados")
        if s.unresolved_metrics:
            add("  METRICAS NO RESUELTAS (insumo de la Fase 1B):")
            for m in s.unresolved_metrics:
                add("    - %s" % m)
    add("")
    add("politica de faltantes: %s" % report.missing_policy)
    add("politica de cuarentena: %s" % report.quarantine_policy)
    add("politica de gitignore: %s" % report.gitignore_policy)
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Protección de datos restringidos (reutilizable)
# ---------------------------------------------------------------------------
def assert_execution_environment_allowed(
    path: Path | None = None,
    env: Mapping[str, str] | None = None,
    repo_root: str | Path | None = None,
) -> None:
    """
    Autoriza -- o no -- ejecutar la inspeccion REAL en este entorno.

    Rechaza dos situaciones, y ninguna es una molestia burocratica:

      1. Codespaces. Los datos de IMARPE tienen restriccion de transferencia y
         Codespaces es una maquina remota de un tercero. Aunque el archivo no
         se versionara, ponerlo alli ya seria transferirlo.
      2. El archivo DENTRO del arbol del repositorio, aunque .gitignore lo
         oculte y aunque Git no lo rastree. Ignorado no es ausente: sigue en el
         disco del proyecto, viaja en cualquier copia del directorio y basta un
         `git add -f` o un cambio de .gitignore para versionarlo.

    La raiz del repositorio es REPO_ROOT, deducida de la ubicacion de este
    modulo y NO del directorio de trabajo. El argumento repo_root existe solo
    para las pruebas: si dependiera del cwd, bastaria lanzar el comando desde
    un subdirectorio para que un Excel guardado en otra carpeta del proyecto
    quedara "fuera" y la compuerta no disparara.

    Las pruebas sinteticas NO pasan por aqui: construyen sus propios libros y
    llaman directamente a las funciones de inspeccion, de modo que pueden
    ejecutarse en Codespaces sin problema. La compuerta protege el camino que
    toca datos reales, no el que toca datos inventados.
    """
    entorno = os.environ if env is None else env
    if str(entorno.get("CODESPACES", "")).strip().lower() == "true":
        raise ExecutionEnvironmentError(
            "Ejecucion bloqueada: CODESPACES=true. Los datos de IMARPE tienen "
            "restriccion de transferencia y no deben abrirse en un entorno "
            "remoto. Ejecute la inspeccion en el equipo local autorizado. "
            "Las pruebas sinteticas si pueden correrse aqui."
        )

    if path is None:
        return

    try:
        archivo = Path(path).resolve()
        raiz = Path(REPO_ROOT if repo_root is None else repo_root).resolve()
    except Exception as exc:
        raise ExecutionEnvironmentError(
            f"No se pudo resolver la ubicacion del archivo: {type(exc).__name__}."
        ) from exc

    if archivo.is_relative_to(raiz):
        raise ExecutionEnvironmentError(
            "Ejecucion bloqueada: el archivo esta DENTRO del arbol del "
            "repositorio. Que .gitignore lo oculte no basta: ignorado no es "
            "ausente, viaja en cualquier copia del directorio y un `git add -f` "
            "lo versionaria. Muevalo fuera del repositorio. (Se omite la ruta a "
            "proposito.)"
        )


def git_tracked_files(repo_root: str | Path | None = None) -> tuple[str, ...]:
    """
    Lista los archivos rastreados por Git.

    La raiz es REPO_ROOT por la misma razon que en la compuerta de entorno:
    lanzada desde un subdirectorio, `git ls-files` devolveria solo esa rama del
    arbol y la comprobacion pareceria limpia sin haber mirado el resto.

    Si Git no esta disponible o el comando falla, se LANZA error. Devolver una
    lista vacia haria que la comprobacion de gobernanza pareciera correcta justo
    cuando no pudo comprobarse nada, que es el peor resultado posible.
    """
    raiz = REPO_ROOT if repo_root is None else repo_root
    try:
        r = subprocess.run(
            ["git", "ls-files"], cwd=str(raiz),
            capture_output=True, text=True, check=False,
        )
    except FileNotFoundError as exc:
        raise GovernanceCheckError(
            "No se encontro el ejecutable de git: la comprobacion de datos "
            "restringidos NO pudo realizarse. No se asume que este limpia."
        ) from exc
    except Exception as exc:
        raise GovernanceCheckError(
            f"No se pudo ejecutar git ls-files: {type(exc).__name__}. La "
            "comprobacion de datos restringidos NO pudo realizarse."
        ) from exc
    if r.returncode != 0:
        raise GovernanceCheckError(
            "git ls-files fallo con codigo %d. La comprobacion de datos "
            "restringidos NO pudo realizarse en %s." % (r.returncode, raiz)
        )
    return tuple(line for line in r.stdout.splitlines() if line.strip())


def find_tracked_restricted_files(
    tracked: Iterable[str],
    suffixes: Sequence[str] = DEFAULT_RESTRICTED_SUFFIXES,
    allowlist: Sequence[str] = (),
) -> tuple[str, ...]:
    """
    Archivos RASTREADOS que portan datos por su extension.

    El criterio es la EXTENSION, no el nombre: un archivo renombrado sigue
    llevando los mismos datos dentro, de modo que exigir que la ruta mencione
    IMARPE dejaria pasar exactamente el caso mas probable de fuga.

    La comparacion es insensible a mayusculas. La allowlist es de rutas exactas
    y esta vacia por defecto: cada excepcion debe ser una decision consciente.
    """
    exts = tuple(str(s).lower() for s in suffixes)
    permitidos = {str(a).strip().lower() for a in allowlist}
    hallados = []
    for ruta in tracked:
        r = str(ruta).strip()
        if not r:
            continue
        bajo = r.lower()
        if bajo.endswith(exts) and bajo not in permitidos:
            hallados.append(r)
    return tuple(sorted(set(hallados)))


def assert_no_restricted_tracked(
    tracked: Iterable[str] | None = None,
    cfg: Mapping | None = None,
    repo_root: str | Path | None = None,
) -> tuple[str, ...]:
    """
    Falla si hay archivos portadores de datos rastreados por Git.

    Es la comprobacion de gobernanza del paquete: el error mas caro de esta
    fase no es un calculo mal hecho, es un archivo restringido que acaba en un
    repositorio y del que ya no se puede volver.
    """
    if cfg is None:
        suffixes: Sequence[str] = DEFAULT_RESTRICTED_SUFFIXES
        allowlist: Sequence[str] = ()
    else:
        prot = cfg.get("proteccion_datos") or {}
        suffixes = tuple(prot.get("extensiones") or DEFAULT_RESTRICTED_SUFFIXES)
        allowlist = tuple(prot.get("allowlist") or ())
    lista = git_tracked_files(repo_root) if tracked is None else tuple(tracked)
    hallados = find_tracked_restricted_files(lista, suffixes, allowlist)
    if hallados:
        raise GovernanceCheckError(
            "FUGA DE DATOS RESTRINGIDOS: estos archivos estan rastreados por Git y "
            "no deberian estarlo: %s. %s" % (list(hallados), GITIGNORE_POLICY)
        )
    return hallados


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------
def main(config_path: str | Path = DEFAULT_CONFIG_PATH) -> int:
    """Ejecuta la inspeccion y escribe el informe. Sin red, sin escritura de datos."""
    try:
        cfg = load_config(config_path)
        assert_execution_environment_allowed()          # entorno, antes de nada
        assert_no_restricted_tracked(cfg=cfg)
        ruta = resolve_source_path(cfg)
        assert_execution_environment_allowed(ruta)      # entorno + ubicacion
        informe = inspect_workbook(ruta, cfg)
    except ImarpeInspectionError as exc:
        print("ERROR | %s | %s" % (type(exc).__name__, exc))
        return 1
    print(format_report(informe))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
